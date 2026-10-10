#!/usr/bin/env python3
"""
Master Report Builder for DCCN Laboratory Experiments 10 to 15
Author: Kanak Sanjay Waradkar (Roll No: 24B-CO-030)
Institution: Goa College of Engineering
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Colors
COLOR_NAVY = "1E3A8A"
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)
COLOR_TEXT = RGBColor(0x11, 0x18, 0x27)
COLOR_MUTED = RGBColor(0x4B, 0x55, 0x63)
COLOR_BORDER = "CBD5E1"
COLOR_BG_LIGHT = "F8FAFC"
COLOR_HEADER_BG = "1E3A8A"

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def setup_header_footer(section, start_page=None):
    # Margins 1 inch
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Header
    header = section.header
    header.is_linked_to_previous = False
    hp1 = header.paragraphs[0]
    hp1.text = ""
    hp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = hp1.add_run("GOA COLLEGE OF ENGINEERING")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    hp2 = header.add_paragraph()
    hp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = hp2.add_run('"Bhausaheb Bandodkar Technical Education Complex"')
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Configure tabs in paragraph properties
    pPr = fp._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    
    # Tab 1: Center at 3.25 inches (4680 dxa)
    tab1 = OxmlElement('w:tab')
    tab1.set(qn('w:val'), 'center')
    tab1.set(qn('w:pos'), '4680')
    tabs.append(tab1)
    
    # Tab 2: Right at 6.5 inches (9360 dxa)
    tab2 = OxmlElement('w:tab')
    tab2.set(qn('w:val'), 'right')
    tab2.set(qn('w:pos'), '9360')
    tabs.append(tab2)
    pPr.append(tabs)

    # Run 1: Roll No
    fr1 = fp.add_run("24B-CO-030\t")
    fr1.font.name = "Times New Roman"
    fr1.font.size = Pt(12)

    # Run 2: Name
    fr2 = fp.add_run("Kanak Sanjay Waradkar\t")
    fr2.font.name = "Times New Roman"
    fr2.font.size = Pt(12)

    # Run 3: Page Field
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    fp._p.append(fldSimple)

    if start_page is not None:
        sectPr = section._sectPr
        pgNumType = OxmlElement('w:pgNumType')
        pgNumType.set(qn('w:start'), str(start_page))
        sectPr.append(pgNumType)

def add_exp_header_block(doc, exp_num, date_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    
    # Configure right tab for date
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), '9360')
    tabs.append(tab)
    pPr.append(tabs)

    r1 = p.add_run(f"Experiment No: {exp_num}")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.bold = True

    r2 = p.add_run(f"\tDate: {date_str}")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.bold = True

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return p

def add_table_data(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows)+1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Headers
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[i], COLOR_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)

    # Data Rows
    for row_idx, data_row in enumerate(rows):
        cells = table.rows[row_idx+1].cells
        bg_col = COLOR_BG_LIGHT if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_data in enumerate(data_row):
            cells[col_idx].text = ""
            p = cells[col_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            # Center if short string or IP, else left
            if len(str(cell_data)) <= 15 or '/' in str(cell_data):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(cell_data))
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            set_cell_background(cells[col_idx], bg_col)
            set_cell_margins(cells[col_idx], top=70, bottom=70, left=100, right=100)

    # Widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Space after table
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(0)
    sp_p.paragraph_format.space_after = Pt(6)

def add_image_figure(doc, img_path, caption, width=Inches(6.0)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        p_img.add_run().add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r = p_cap.add_run(caption)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        r.font.italic = True
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

def add_screenshot_box(doc, caption_text, height=Inches(2.5)):
    # Create single cell bordered table with shaded background
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    
    # Border & Shading
    set_cell_background(cell, "F8FAFC")
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'  <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'  <w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'  <w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(25)
    p1.paragraph_format.space_after = Pt(4)
    r1 = p1.add_run("[ SCREENSHOT PLACEHOLDER ]")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(25)
    r2 = p2.add_run(caption_text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(8)

def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_str.strip())
    r.font.name = "Courier New"
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

print("Builder helper functions compiled successfully!")
